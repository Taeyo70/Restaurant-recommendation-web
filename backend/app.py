from flask import Flask, render_template, request, redirect, url_for, session
import sqlite3
import json
import tempfile
import time
from threading import Lock
import os
import urllib.request
import urllib.error

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash, check_password_hash


app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-me")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.path.join(BASE_DIR, "restaurants.db")

AUTH_WORKBOOK = os.environ.get(
    "RESTAURANT_WORKBOOK",
    os.path.abspath(
        os.path.join(
            BASE_DIR,
            "..",
            "data",
            "restaurants_final_with_maps.xlsx"
        )
    ),
)

AUTH_LOCK = Lock()

# ============================================================
# AI / OPENROUTER
# ============================================================

OPENROUTER_API_KEY = os.environ.get("OPENROUTER_API_KEY")

print("OPENROUTER KEY LOADED:", bool(OPENROUTER_API_KEY))

OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"

AI_MODEL = "openrouter/free"


def ask_ai(prompt, temperature=0.2, max_tokens=700):
    """
    Send a prompt to OpenRouter and return the AI text response.
    """

    if not OPENROUTER_API_KEY:
        raise RuntimeError(
            "OPENROUTER_API_KEY is not configured."
        )

    payload = {
        "model": AI_MODEL,
        "messages": [
            {
                "role": "user",
                "content": prompt
            }
        ],
        "temperature": temperature,
        "max_tokens": max_tokens
    }

    data = json.dumps(payload).encode("utf-8")

    req = urllib.request.Request(
        OPENROUTER_URL,
        data=data,
        headers={
            "Authorization": f"Bearer {OPENROUTER_API_KEY}",
            "Content-Type": "application/json",
            "HTTP-Referer": "https://restaurant-finder.onrender.com",
            "X-Title": "Restaurant Finder",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            response_data = json.loads(
                response.read().decode("utf-8")
            )

    except urllib.error.HTTPError as error:
        error_body = error.read().decode("utf-8", errors="ignore")
        print("OpenRouter HTTP error:", error.code, error_body)
        raise RuntimeError(
            f"OpenRouter HTTP error {error.code}"
        )

    except urllib.error.URLError as error:
        print("OpenRouter connection error:", error)
        raise RuntimeError(
            "Could not connect to OpenRouter."
        )

    except Exception as error:
        print("OpenRouter error:", error)
        raise

    try:
        return response_data["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError):
        print("Unexpected OpenRouter response:", response_data)
        raise RuntimeError(
            "Invalid response received from OpenRouter."
        )


# ============================================================
# DATABASE
# ============================================================

def get_db():
    conn = sqlite3.connect(
        DATABASE,
        timeout=10
    )

    conn.row_factory = sqlite3.Row

    return conn


# ============================================================
# AUTH WORKBOOK
# ============================================================

def _open_auth_workbook():

    if not os.path.exists(AUTH_WORKBOOK):
        raise FileNotFoundError(
            f"Workbook not found: {AUTH_WORKBOOK}"
        )

    workbook = load_workbook(AUTH_WORKBOOK)

    if "users" not in workbook.sheetnames:

        users = workbook.create_sheet("users")

        users.append([
            "id",
            "username",
            "email",
            "password_hash",
            "created_at"
        ])

    return workbook


def _save_auth_workbook(workbook):

    directory = (
        os.path.dirname(AUTH_WORKBOOK)
        or "."
    )

    file_descriptor, temporary_path = tempfile.mkstemp(
        dir=directory,
        prefix="restaurant-finder-auth-",
        suffix=".xlsx",
    )

    os.close(file_descriptor)

    try:

        workbook.save(temporary_path)
        workbook.close()

        os.replace(
            temporary_path,
            AUTH_WORKBOOK
        )

    finally:

        if os.path.exists(temporary_path):
            os.remove(temporary_path)


def init_auth_sheets():

    with AUTH_LOCK:

        if not os.path.exists(AUTH_WORKBOOK):
            raise FileNotFoundError(
                f"Workbook not found: {AUTH_WORKBOOK}"
            )

        workbook = load_workbook(
            AUTH_WORKBOOK
        )

        if "users" in workbook.sheetnames:

            workbook.close()
            return

        users = workbook.create_sheet(
            "users"
        )

        users.append([
            "id",
            "username",
            "email",
            "password_hash",
            "created_at"
        ])

        _save_auth_workbook(workbook)


def find_user(identifier):

    with AUTH_LOCK:

        workbook = _open_auth_workbook()

        users = workbook["users"]

        user = None

        for row in users.iter_rows(
            min_row=2,
            values_only=True
        ):

            if (
                row[1] == identifier
                or str(row[2]).lower()
                == identifier.lower()
            ):

                user = {
                    "id": row[0],
                    "username": row[1],
                    "email": row[2],
                    "password_hash": row[3],
                }

                break

        workbook.close()

        return user


def find_user_by_email(email):
    return find_user(email)


init_auth_sheets()


# ============================================================
# LOGIN
# ============================================================

@app.route("/")
def login():

    return render_template(
        "login.html"
    )


@app.route("/login", methods=["POST"])
def do_login():

    identifier = request.form.get(
        "identifier",
        ""
    ).strip()

    password = request.form.get(
        "password",
        ""
    )

    user = find_user(identifier)

    if (
        user is None
        or not check_password_hash(
            user["password_hash"],
            password
        )
    ):

        return render_template(
            "login.html",
            error="Invalid username/email or password."
        )

    session["user_id"] = user["id"]
    session["username"] = user["username"]

    return redirect(
        url_for("home")
    )


# ============================================================
# PROFILE
# ============================================================

@app.route("/profile")
def profile():

    if not session.get("user_id"):
        return redirect(
            url_for("login")
        )

    user = find_user(
        session.get(
            "username",
            ""
        )
    )

    if user is None:

        session.clear()

        return redirect(
            url_for("login")
        )

    return render_template(
        "profile.html",
        user=user
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# ============================================================
# FAVORITES
# ============================================================

@app.route("/favorites")
def favorites():

    return render_template(
        "favorites.html"
    )


# ============================================================
# REGISTER
# ============================================================

@app.route(
    "/register",
    methods=["GET", "POST"]
)
def register():

    if request.method == "GET":

        return render_template(
            "register.html"
        )

    username = request.form.get(
        "username",
        ""
    ).strip()

    email = request.form.get(
        "email",
        ""
    ).strip().lower()

    password = request.form.get(
        "password",
        ""
    )

    if len(username) < 3:

        return render_template(
            "register.html",
            error="Username must contain at least 3 characters."
        )

    if len(password) < 8:

        return render_template(
            "register.html",
            error="Password must contain at least 8 characters."
        )

    password_hash = generate_password_hash(
        password
    )

    with AUTH_LOCK:

        workbook = _open_auth_workbook()

        users = workbook["users"]

        for row in users.iter_rows(
            min_row=2,
            values_only=True
        ):

            if (
                row[1] == username
                or str(row[2]).lower()
                == email
            ):

                workbook.close()

                return render_template(
                    "register.html",
                    error="Username or email already exists."
                )

        user_ids = [
            row[0]
            for row in users.iter_rows(
                min_row=2,
                values_only=True
            )
            if row[0]
        ]

        next_id = (
            max(user_ids, default=0)
            + 1
        )

        users.append([
            next_id,
            username,
            email,
            password_hash,
            int(time.time())
        ])

        _save_auth_workbook(
            workbook
        )

        workbook.close()

    return redirect(
        url_for("login")
    )


# ============================================================
# FORGOT PASSWORD
# ============================================================

@app.route(
    "/forgot-password",
    methods=["GET", "POST"]
)
def forgot_password():

    if request.method == "GET":

        return render_template(
            "forgot_password.html"
        )

    email = request.form.get(
        "email",
        ""
    ).strip().lower()

    new_password = request.form.get(
        "password",
        ""
    )

    if len(new_password) < 8:

        return render_template(
            "forgot_password.html",
            error="Password must contain at least 8 characters."
        )

    password_hash = generate_password_hash(
        new_password
    )

    updated = False

    with AUTH_LOCK:

        workbook = _open_auth_workbook()

        users = workbook["users"]

        for row in users.iter_rows(
            min_row=2
        ):

            if (
                str(row[2].value).lower()
                == email
            ):

                row[3].value = password_hash

                updated = True

                break

        _save_auth_workbook(
            workbook
        )

        workbook.close()

    if not updated:

        return render_template(
            "forgot_password.html",
            error="No account was found with that email address."
        )

    return render_template(
        "forgot_password.html",
        message="Your password was changed. You can now log in."
    )


# ============================================================
# HOME
# ============================================================

@app.route("/home")
def home():

    conn = get_db()

    total_restaurants = conn.execute(
        "SELECT COUNT(*) FROM restaurants"
    ).fetchone()[0]

    total_cities = conn.execute(
        "SELECT COUNT(DISTINCT City) FROM restaurants"
    ).fetchone()[0]

    avg_rating = conn.execute(
        "SELECT ROUND(AVG(Rating), 1) FROM restaurants"
    ).fetchone()[0]

    avg_cost = conn.execute(
        "SELECT ROUND(AVG(Cost), 0) FROM restaurants"
    ).fetchone()[0]

    city_data = conn.execute("""
        SELECT
            City,
            COUNT(*) AS total,
            ROUND(AVG(Rating), 2) AS rating
        FROM restaurants
        WHERE City IS NOT NULL
        GROUP BY City
        ORDER BY total DESC
        LIMIT 8
    """).fetchall()

    city_data = [
        dict(row)
        for row in city_data
    ]

    if city_data:

        highest_city_total = city_data[0]["total"]

        for city in city_data:

            city["percentage"] = round(
                (
                    city["total"]
                    / highest_city_total
                ) * 100
            )

    cuisine_data = conn.execute("""
        SELECT
            Cuisine,
            COUNT(*) AS total
        FROM restaurants
        WHERE Cuisine IS NOT NULL
        GROUP BY Cuisine
        ORDER BY total DESC
        LIMIT 8
    """).fetchall()

    cuisine_data = [
        dict(row)
        for row in cuisine_data
    ]

    if cuisine_data:

        highest_cuisine_total = (
            cuisine_data[0]["total"]
        )

        for cuisine in cuisine_data:

            cuisine["percentage"] = round(
                (
                    cuisine["total"]
                    / highest_cuisine_total
                ) * 100
            )

    conn.close()

    return render_template(
        "home.html",
        total_restaurants=total_restaurants,
        total_cities=total_cities,
        avg_rating=avg_rating,
        avg_cost=avg_cost,
        city_data=city_data,
        cuisine_data=cuisine_data
    )


# ============================================================
# NORMAL SEARCH
# ============================================================


# ============================================================
# NATURAL-LANGUAGE HELPERS
# ============================================================

import re

STOP_WORDS = {
    "a", "an", "the", "in", "at", "near", "around", "for", "to", "of",
    "and", "or", "with", "me", "my", "i", "want", "need", "find", "show",
    "give", "suggest", "looking", "look", "good", "best", "nice", "some",
    "any", "place", "places", "restaurant", "restaurants", "food",
    "eat", "eating", "please", "can", "you", "is", "are", "where",
    "cheap", "budget", "affordable", "top", "rated", "highly",
    "rating", "ratings", "star", "stars", "above", "under", "below",
    "cost", "price", "rs", "rupees", "food", "foods", "items",
}

# plain words people use -> words that actually appear in the data
SYNONYMS = {
    
}


def extract_hints(text):
    """Read 'under 500', 'below 1000', '4+ rating', 'cheap', 'top rated'."""
    max_cost = None
    min_rating = None

    m = re.search(r"(?:under|below|less than|within|max|upto|up to)\s*(?:rs\.?|₹)?\s*(\d{2,5})", text)
    if m:
        max_cost = float(m.group(1))
        text = text.replace(m.group(0), " ")

    m = re.search(r"(\d(?:\.\d)?)\s*(?:\+|plus|star|stars|and above|and up)", text)
    if m:
        min_rating = float(m.group(1))
        text = text.replace(m.group(0), " ")

    if re.search(r"\b(cheap|budget|affordable|inexpensive)\b", text) and not max_cost:
        max_cost = 600.0

    if re.search(r"\b(top rated|highly rated|best)\b", text) and not min_rating:
        min_rating = 4.0

    return {"text": text, "max_cost": max_cost, "min_rating": min_rating}


def meaningful_words(text):
    words = re.findall(r"[a-z]+", text.lower())
    out = []
    for w in words:
        if w in STOP_WORDS or len(w) < 3:
            continue
        w = SYNONYMS.get(w, w)
        out.append(w)
    return out



# ============================================================
# CUISINE DETECTION (uses every cuisine that exists in the DB)
# ============================================================

_CUISINE_CACHE = {"list": None}

CUISINE_ALIASES = {
    "dessert": "Desserts", "sweet": "Desserts", "sweets": "Desserts",
    "cake": "Bakery and Confectionary", "bakery": "Bakery and Confectionary",
    "veg": "North Indian", "maharashtra": "Maharashtrian",
    "south": "South Indian", "dosa": "South Indian", "idli": "South Indian",
    "biriyani": "Biryani", "noodles": "Chinese", "momos": "Tibetan",
    "bbq": "Barbecue", "cafe": "Coffee", "chai": "Tea",
    "mexican": "Mexican", "arab": "Arabian",
}


def all_cuisines(conn):
    if _CUISINE_CACHE["list"] is None:
        found = set()
        for row in conn.execute(
            "SELECT DISTINCT Cuisine FROM restaurants WHERE Cuisine IS NOT NULL"
        ).fetchall():
            for item in row["Cuisine"].split(","):
                item = item.strip()
                if item:
                    found.add(item)
        # longest first so "north indian" wins over "indian"
        _CUISINE_CACHE["list"] = sorted(found, key=len, reverse=True)
    return _CUISINE_CACHE["list"]


def find_cuisines(text, conn):
    """Return (cuisines_found, text_with_cuisine_words_removed)."""
    text = text.lower()
    found = []
    for cuisine in all_cuisines(conn):
        base = cuisine.lower()
        if base.endswith("s"):
            base = base[:-1]
        pattern = r"\b" + re.escape(base) + r"s?\b"
        if re.search(pattern, text):
            found.append(cuisine)
            text = re.sub(pattern, " ", text)
    for alias, cuisine in CUISINE_ALIASES.items():
        pattern = r"\b" + re.escape(alias) + r"s?\b"
        if re.search(pattern, text):
            if cuisine not in found:
                found.append(cuisine)
            text = re.sub(pattern, " ", text)
    return found, text



def _norm(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


def parse_ai_reasons(ai_text, restaurants):
    """Pull {restaurant name: reason} out of messy AI output."""
    ai_text = (ai_text or "").strip()
    start, end = ai_text.find("{"), ai_text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("No JSON found in AI reply")
    data = json.loads(ai_text[start:end + 1])
    items = data.get("reasons", []) if isinstance(data, dict) else data

    by_norm = {_norm(r["Name"]): r["Name"] for r in restaurants}
    result = {}
    for pos, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        reason = (item.get("reason") or "").strip()
        if not reason:
            continue
        name = None
        key = _norm(item.get("name", ""))
        if key in by_norm:
            name = by_norm[key]
        else:  # partial match, e.g. AI dropped the branch name
            for k, original in by_norm.items():
                if key and (key in k or k in key):
                    name = original
                    break
        if name is None and pos < len(restaurants):
            name = restaurants[pos]["Name"]  # fall back to order
        if name and name not in result:
            result[name] = reason
    return result


# ============================================================
# HUMAN-SOUNDING RECOMMENDATION REASONS (used when AI is off)
# ============================================================

import random


def friendly_reason(r, prefs, used=None):
    """Build a specific, varied reason from the restaurant's real data."""
    name = r.get("Name", "This place")
    rating = r.get("Rating") or 0
    votes = r.get("Votes") or 0
    cost = r.get("Cost") or 0
    cuisines = [c.strip() for c in (r.get("Cuisine") or "").split(",") if c.strip()]
    wanted = (prefs.get("cuisine") or "").lower()
    main = next((c for c in cuisines if wanted and wanted in c.lower()),
                cuisines[0] if cuisines else "food")
    area = (r.get("Location") or r.get("City") or "").split(",")[0].strip()

    picks = []
    if rating >= 4.3:
        picks.append(f"Okay this one's a gem — {rating}★ for {main}, you can't really go wrong.")
        picks.append(f"{rating}★ is hard to beat. If you want {main}, this is your safe bet.")
        picks.append(f"Top of the list for a reason: {rating}★ and great {main}.")
    elif rating >= 3.8:
        picks.append(f"Solid {rating}★ for {main}, and people keep coming back for more.")
        picks.append(f"Reliable pick — {rating}★ and a proper {main} menu.")
        picks.append(f"Good {main} spot — {rating}★ from real diners.")
    else:
        picks.append(f"A decent {main} option at {rating}★ if you want something easy and nearby.")

    if votes >= 500:
        picks.append(f"Loved by a crowd — {int(votes)}+ votes, so it's a proven {main} place.")
    if cost and cost <= 500:
        picks.append(f"Easy on the pocket at ~₹{int(cost)} for two, and {rating}★ {main} too.")
    elif cost and cost >= 1500:
        picks.append(f"A bit of a splurge (~₹{int(cost)} for two), but {name} is worth it for {main}.")
    elif cost:
        picks.append(f"Fair price (~₹{int(cost)} for two) for proper {main} at {name}.")
    if area:
        picks.append(f"It's in {area}, so you get your {main} fix without a long trip.")
    if prefs.get("mood"):
        picks.append(f"Perfect for a {prefs['mood'].lower()} kind of day — {main} hits the spot.")
    if prefs.get("meal"):
        picks.append(f"Great for {prefs['meal'].lower()} — {main} done right, {rating}★.")
    if len(cuisines) >= 3:
        picks.append(f"Lots of choice here — {', '.join(cuisines[:3])} all on one menu.")

    random.seed(name)  # same restaurant -> same line, different restaurants vary
    random.shuffle(picks)
    used = used or set()
    for p in picks:
        if p not in used:
            return p
    return f"{name}: {picks[0]}"  # every line already used -> make it unique


@app.route("/search")
def search():

    keyword = request.args.get(
        "keyword",
        ""
    ).strip()

    city = request.args.get(
        "city",
        ""
    ).strip()

    cuisine = request.args.get(
        "cuisine",
        ""
    ).strip()

    max_cost = request.args.get(
        "max_cost",
        ""
    ).strip()

    min_rating = request.args.get(
        "min_rating",
        ""
    ).strip()

    conn = get_db()

    query = """
        SELECT *
        FROM restaurants
        WHERE 1=1
    """

    params = []

    if keyword:

        text = keyword.lower()

        # 1) Detect a city mentioned anywhere in the sentence
        cities = conn.execute(
            "SELECT DISTINCT City FROM restaurants"
        ).fetchall()

        found_city = None

        for row in cities:
            if row["City"] and row["City"].lower() in text:
                found_city = row["City"]
                text = text.replace(row["City"].lower(), " ")
                break

        if found_city and not city:
            query += " AND City = ? "
            params.append(found_city)

        # 1b) Detect cuisines (north indian, malvani, chinese, desserts...)
        found_cuisines, text = find_cuisines(text, conn)

        if found_cuisines:
            query += " AND (" + " OR ".join(
                ["Cuisine LIKE ?"] * len(found_cuisines)
            ) + ") "
            params.extend(["%" + c + "%" for c in found_cuisines])

        # 2) Pull out price / rating hints written in plain language
        hints = extract_hints(text)
        text = hints["text"]

        if hints["max_cost"] and not max_cost:
            query += " AND Cost <= ? "
            params.append(hints["max_cost"])

        if hints["min_rating"] and not min_rating:
            query += " AND Rating >= ? "
            params.append(hints["min_rating"])

        # 3) Every remaining meaningful word must match somewhere
        #    (word-by-word, NOT the whole sentence as one phrase)
        for word in meaningful_words(text):
            query += """
                AND (
                    Name LIKE ?
                    OR Cuisine LIKE ?
                    OR Location LIKE ?
                    OR Locality LIKE ?
                )
            """
            value = "%" + word + "%"
            params.extend([value, value, value, value])

    if city:

        query += """
            AND City = ?
        """

        params.append(city)

    if cuisine:

        query += """
            AND Cuisine LIKE ?
        """

        params.append(
            "%"
            + cuisine
            + "%"
        )

    if max_cost:

        try:

            query += """
                AND Cost <= ?
            """

            params.append(
                float(max_cost)
            )

        except ValueError:
            pass

    if min_rating:

        try:

            query += """
                AND Rating >= ?
            """

            params.append(
                float(min_rating)
            )

        except ValueError:
            pass

    query += """
        ORDER BY Rating DESC
        LIMIT 60
    """

    restaurants = conn.execute(
        query,
        params
    ).fetchall()

    cities = conn.execute("""
        SELECT DISTINCT City
        FROM restaurants
        WHERE City IS NOT NULL
        AND City != ''
        ORDER BY City
    """).fetchall()

    cuisines_raw = conn.execute("""
        SELECT DISTINCT Cuisine
        FROM restaurants
        WHERE Cuisine IS NOT NULL
        AND Cuisine != ''
    """).fetchall()

    cuisine_set = set()

    for row in cuisines_raw:

        if row["Cuisine"]:

            for item in row["Cuisine"].split(","):

                c = item.strip()

                if c:
                    cuisine_set.add(c)

    sorted_cuisines = sorted(
        list(cuisine_set)
    )

    conn.close()

    return render_template(
        "search.html",
        restaurants=restaurants,
        cities=cities,
        cuisines=sorted_cuisines,
        selected_city=city,
        selected_cuisine=cuisine,
        selected_cost=max_cost,
        selected_rating=min_rating,
        keyword=keyword
    )


# ============================================================
# AI SEARCH
# ============================================================

@app.route(
    "/ai-search",
    methods=["POST"]
)
def ai_search():

    user_query = request.form.get(
        "query",
        ""
    ).strip()

    if not user_query:

        return redirect(
            url_for("search")
        )

    conn = get_db()
    valid_cities = [r["City"] for r in conn.execute(
        "SELECT DISTINCT City FROM restaurants WHERE City IS NOT NULL"
    ).fetchall()]
    valid_cuisines = sorted(all_cuisines(conn))
    conn.close()

    prompt = f"""
You convert a person's natural-language restaurant request into search filters.

Request:
{user_query}

Valid cities (use EXACT spelling, or "" if none is mentioned):
{", ".join(valid_cities)}

All cuisines in the database (use EXACT spelling):
{", ".join(valid_cuisines)}

Return ONLY valid JSON with exactly these fields:

{{
    "keyword": "",
    "city": "",
    "cuisine": "",
    "max_cost": "",
    "min_rating": ""
}}

Rules:
- city must be one of the valid cities above, otherwise "".
- cuisine: pick ONE cuisine from the list above if the request implies it
  (e.g. "malvani fish" -> "Malvani", "sweet" -> "Desserts",
  "noodles" -> "Chinese"). Must be an exact match from the list, else "".
- keyword: only a specific dish or restaurant name, otherwise "".
  Never put the city or the whole sentence in keyword.
- max_cost: number for cost for two (cheap = 500, mid-range = 1000).
- min_rating: number 1-5 (good = 4, very good / best = 4.3).
- Use "" when information is missing.
- No explanation, no markdown, JSON only.
"""

    try:

        ai_text = ask_ai(
            prompt
        )

        # Remove possible markdown fences
        ai_text = ai_text.strip()

        if ai_text.startswith("```"):

            ai_text = (
                ai_text
                .replace("```json", "")
                .replace("```", "")
                .strip()
            )

        filters = json.loads(
            ai_text
        )

    except Exception as error:

        print(
            "AI SEARCH ERROR:",
            error
        )

        return redirect(
            url_for(
                "search",
                keyword=user_query
            )
        )

    allowed_filters = [
        "keyword",
        "city",
        "cuisine",
        "max_cost",
        "min_rating"
    ]

    cleaned_filters = {
        key: filters.get(
            key,
            ""
        )
        for key in allowed_filters
        if filters.get(
            key,
            ""
        ) not in ("", None)
    }

    # Validate AI cuisine/city against the real database values
    if "cuisine" in cleaned_filters:
        match = [c for c in valid_cuisines
                 if c.lower() == str(cleaned_filters["cuisine"]).lower()]
        if match:
            cleaned_filters["cuisine"] = match[0]
        else:
            cleaned_filters.pop("cuisine")

    if "city" in cleaned_filters:
        match = [c for c in valid_cities
                 if c.lower() == str(cleaned_filters["city"]).lower()]
        if match:
            cleaned_filters["city"] = match[0]
        else:
            cleaned_filters.pop("city")

    # Nothing usable from the AI -> let the smart keyword search handle it
    if not cleaned_filters:
        return redirect(url_for("search", keyword=user_query))

    return redirect(
        url_for(
            "search",
            **cleaned_filters
        )
    )


# ============================================================
# RECOMMENDATION KEYWORDS
# ============================================================

MOOD_KEYWORDS = {

    "Comfort food": [
        "Biryani",
        "North Indian",
        "Fast Food",
        "Chinese"
    ],

    "Date night": [
        "Italian",
        "Continental",
        "Cafe"
    ],

    "Healthy": [
        "Healthy",
        "Salad",
        "Vegan",
        "Vegetarian"
    ],

    "Spicy": [
        "Indian",
        "Mexican",
        "Thai",
        "Biryani"
    ],

    "Quick bite": [
        "Fast Food",
        "Burger",
        "Pizza",
        "Street Food"
    ],

    "Sweet craving": [
        "Dessert",
        "Bakery",
        "Ice Cream",
        "Cafe"
    ],
}


MEAL_KEYWORDS = {

    "Breakfast": [
        "Breakfast",
        "Cafe",
        "Bakery"
    ],

    "Snacks": [
        "Fast Food",
        "Street Food",
        "Cafe"
    ],

    "Dinner": [
        "Dinner",
        "Fine Dining",
        "North Indian",
        "Chinese"
    ],
}


DIET_KEYWORDS = {

    "Vegetarian": [
        "Vegetarian",
        "South Indian",
        "North Indian",
        "Jain",
        "Pure Veg"
    ],

    "Vegan": [
        "Vegan",
        "Salad",
        "Health Food",
        "Juice"
    ],

    "Non-veg": [
        "Mughlai",
        "Kebab",
        "Tandoori",
        "Seafood",
        "Biryani",
        "Chicken",
        "Mutton",
        "Fish"
    ],
}


# ============================================================
# SMART RECOMMENDATIONS
# ============================================================

@app.route(
    "/recommendations",
    methods=["GET", "POST"]
)
@app.route(
    "/recommend",
    methods=["POST"]
)
def recommendations():

    preferences = {

        "city": request.form.get(
            "city",
            ""
        ).strip(),

        "mood": request.form.get(
            "mood",
            ""
        ).strip(),

        "meal": request.form.get(
            "meal",
            ""
        ).strip(),

        "cuisine": request.form.get(
            "cuisine",
            ""
        ).strip(),

        "diet": request.form.get(
            "diet",
            ""
        ).strip(),

        "max_cost": request.form.get(
            "max_cost",
            ""
        ).strip(),

        "min_rating": request.form.get(
            "min_rating",
            ""
        ).strip(),
    }

    if request.method == "GET":

        return render_template(
            "recommendations.html",
            preferences=preferences,
            restaurants=[],
            searched=False,
            ai_message=""
        )

    session["preferences"] = preferences

    conditions = [
        "1=1"
    ]

    params = []

    # --------------------------------------------------------
    # CITY
    # --------------------------------------------------------

    if preferences["city"]:

        conditions.append(
            "City LIKE ?"
        )

        params.append(
            "%"
            + preferences["city"]
            + "%"
        )

    # --------------------------------------------------------
    # CUISINE
    # --------------------------------------------------------

    if preferences["cuisine"]:

        conditions.append(
            "Cuisine LIKE ?"
        )

        params.append(
            "%"
            + preferences["cuisine"]
            + "%"
        )

    # Mood / meal / diet are SOFT preferences: if they leave nothing,
    # we retry without them instead of showing an empty page.
    soft_conditions = []
    soft_params = []
    hard_conditions, hard_params = conditions, params
    conditions, params = soft_conditions, soft_params

    # --------------------------------------------------------
    # MOOD + MEAL
    # --------------------------------------------------------

    terms = []

    terms += MOOD_KEYWORDS.get(
        preferences["mood"],
        []
    )

    terms += MEAL_KEYWORDS.get(
        preferences["meal"],
        []
    )

    terms = list(
        dict.fromkeys(terms)
    )

    if terms:

        term_conditions = []

        for term in terms:

            term_conditions.append(
                """
                (
                    Cuisine LIKE ?
                    OR Name LIKE ?
                )
                """
            )

            params.extend([
                "%"
                + term
                + "%",

                "%"
                + term
                + "%"
            ])

        conditions.append(
            "("
            + " OR ".join(
                term_conditions
            )
            + ")"
        )

    # --------------------------------------------------------
    # DIET
    # --------------------------------------------------------

    if (
        preferences["diet"]
        and preferences["diet"]
        != "No preference"
    ):

        diet_terms = DIET_KEYWORDS.get(
            preferences["diet"],
            [
                preferences["diet"]
            ]
        )

        diet_conditions = []

        for term in diet_terms:

            diet_conditions.append(
                """
                (
                    Cuisine LIKE ?
                    OR Name LIKE ?
                )
                """
            )

            params.extend([
                "%"
                + term
                + "%",

                "%"
                + term
                + "%"
            ])

        conditions.append(
            "("
            + " OR ".join(
                diet_conditions
            )
            + ")"
        )

    conditions, params = hard_conditions, hard_params

    # --------------------------------------------------------
    # MAX COST
    # --------------------------------------------------------

    try:

        if preferences["max_cost"]:

            conditions.append(
                "Cost <= ?"
            )

            params.append(
                float(
                    preferences["max_cost"]
                )
            )

    except ValueError:

        preferences["max_cost"] = ""

    # --------------------------------------------------------
    # MIN RATING
    # --------------------------------------------------------

    try:

        if preferences["min_rating"]:

            conditions.append(
                "Rating >= ?"
            )

            params.append(
                float(
                    preferences["min_rating"]
                )
            )

    except ValueError:

        preferences["min_rating"] = ""

    # --------------------------------------------------------
    # DATABASE QUERY
    # --------------------------------------------------------

    base_sql = """
        SELECT
            Name,
            Location,
            City,
            Cuisine,
            Rating,
            Votes,
            Cost,
            "Google Maps Link"
        FROM restaurants
        WHERE {where}
        ORDER BY Rating DESC, Votes DESC
        LIMIT 12
    """

    conn = get_db()

    relaxed = False

    restaurants = [
        dict(row)
        for row in conn.execute(
            base_sql.format(
                where=" AND ".join(hard_conditions + soft_conditions)
            ),
            hard_params + soft_params
        ).fetchall()
    ]

    if not restaurants and soft_conditions:
        relaxed = True
        restaurants = [
            dict(row)
            for row in conn.execute(
                base_sql.format(where=" AND ".join(hard_conditions)),
                hard_params
            ).fetchall()
        ]

    conn.close()

    # --------------------------------------------------------
    # AI EXPLANATIONS
    # --------------------------------------------------------

    reasons = {}

    if not restaurants:

        ai_message = (
            "No matches found. "
            "Try changing your preferences."
        )

    elif not OPENROUTER_API_KEY:

        ai_message = (
            "✨ Smart picks based on ratings, price and your preferences."
        )

    else:

        slim = [
            {
                "name": r["Name"],
                "cuisine": r["Cuisine"],
                "rating": r["Rating"],
                "votes": r["Votes"],
                "cost_for_two": r["Cost"],
                "area": r["Location"],
                "city": r["City"],
            }
            for r in restaurants
        ]

        prompt = f"""
You are a friendly foodie friend recommending restaurants.

What the person wants:
{json.dumps({k: v for k, v in preferences.items() if v}, indent=2)}

Restaurants:
{json.dumps(slim, indent=2)}

For EACH restaurant write ONE short, warm, casual line (12-22 words)
saying why it suits THIS person. Sound like a friend texting, not a robot.

Rules:
- Mention something specific: its rating, price, cuisine, area or votes.
- Connect it to what they asked for (mood, meal, cuisine, budget, diet).
- Every line must be different. Never repeat the same opening words.
- Never write "matches your preferences" or "matches your selected".
- Use only the facts given. Do not invent dishes or reviews.
- Return ONLY valid JSON, no markdown:

{{"reasons": [{{"name": "exact restaurant name", "reason": "the line"}}]}}
"""

        try:

            reasons = {}
            last_error = None

            for attempt in range(2):  # free models fail often: retry once
                try:
                    ai_text = ask_ai(
                        prompt,
                        temperature=0.8,
                        max_tokens=2500
                    )
                    reasons = parse_ai_reasons(ai_text, restaurants)
                    if reasons:
                        break
                except Exception as err:
                    last_error = err
                    print("RECOMMENDATION AI attempt", attempt + 1, "failed:", err)

            if reasons:
                ai_message = "✨ AI personalized these recommendations for you."
            else:
                ai_message = "✨ Smart picks based on ratings, price and your preferences."
                if last_error:
                    print("RECOMMENDATION AI gave up:", last_error)

        except Exception as error:

            print(
                "RECOMMENDATION AI ERROR:",
                error
            )

            reasons = {}

            ai_message = (
                "Showing database matches. "
                
            )

    # --------------------------------------------------------
    # ADD REASONS TO RESTAURANT CARDS
    # --------------------------------------------------------

    used = set()

    for restaurant in restaurants:

        reason = reasons.get(restaurant["Name"])

        # missing or repeated -> build a specific one from the data
        if not reason or reason in used:
            reason = friendly_reason(restaurant, preferences, used)

        used.add(reason)
        restaurant["reason"] = reason

    if relaxed and restaurants:
        ai_message = (
            "No exact match for your mood/diet, so here are the "
            "closest options for your city, cuisine and budget."
        )

    # --------------------------------------------------------
    # RENDER PAGE
    # --------------------------------------------------------

    return render_template(
        "recommendations.html",
        preferences=preferences,
        restaurants=restaurants,
        searched=True,
        ai_message=ai_message,
    )


# ============================================================
# RUN APPLICATION
# ============================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=True
    )